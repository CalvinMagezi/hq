"""Generate the synthetic conversion fixtures. All names, numbers and amounts are invented.

Usage: python3 generate.py <out dir> [statement rows]
Requires python-docx, openpyxl, python-pptx and Pillow. See README.md for the
LibreOffice steps that produce the other formats.
"""
import os, random, datetime, struct, zlib, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt as PPt
from PIL import Image, ImageDraw, ImageFont, ImageFilter

OUT = sys.argv[1]
ROWS = int(sys.argv[2]) if len(sys.argv) > 2 else 52
os.makedirs(OUT, exist_ok=True)
random.seed(7)

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    st = OxmlElement("w:rStyle"); st.set(qn("w:val"), "Hyperlink"); rpr.append(st)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; r.append(t); h.append(r)
    paragraph._p.append(h)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color); tcPr.append(shd)

def png(path, w, h, color):
    raw = b"".join(b"\x00" + bytes(color) * w for _ in range(h))
    def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    open(path, "wb").write(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))

# ---------- invoice.docx ----------
d = Document()
d.add_heading("INVOICE", 0)
p = d.add_paragraph(); r = p.add_run("Example Agro Supplies Ltd"); r.bold = True
d.add_paragraph("Plot 12, Example Road, Kampala, Uganda\nTel: 0700 000 000")
t = d.add_table(rows=2, cols=2); t.style = "Table Grid"
t.cell(0, 0).text = "Bill to"; t.cell(0, 1).text = "Invoice details"
t.cell(1, 0).text = "Sample Farmers Cooperative\nMukono District"; t.cell(1, 1).text = "Invoice No: INV-0042\nDate: 05 October 2026\nDue: 19 October 2026"
for c in t.rows[0].cells:
    shade(c, "DDEEDD"); c.paragraphs[0].runs[0].bold = True
d.add_paragraph()
items = [("Maize seed, 25 kg bag", 10, 85000), ("NPK fertiliser, 50 kg", 6, 160000), ("Transport to Mukono", 1, 120000), ("Field training day", 2, 250000)]
t = d.add_table(rows=1, cols=4); t.style = "Table Grid"
for i, h in enumerate(["Description", "Qty", "Unit price (UGX)", "Amount (UGX)"]):
    t.rows[0].cells[i].text = h; t.rows[0].cells[i].paragraphs[0].runs[0].bold = True; shade(t.rows[0].cells[i], "DDEEDD")
total = 0
for desc, q, u in items:
    row = t.add_row().cells; amt = q * u; total += amt
    row[0].text = desc; row[1].text = str(q); row[2].text = f"{u:,}"; row[3].text = f"{amt:,}"
    for c in row[1:]: c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
row = t.add_row().cells; row[0].text = "TOTAL"; row[3].text = f"{total:,}"
row[0].paragraphs[0].runs[0].bold = True; row[3].paragraphs[0].runs[0].bold = True
d.add_paragraph()
d.add_heading("Payment details", 2)
for s in ["Bank: Example Bank, account 0000000000", "Mobile Money: 0700 000 000 (Example Agro Supplies)", "Reference: INV-0042"]:
    d.add_paragraph(s, style="List Bullet")
p = d.add_paragraph("Pay online at "); add_hyperlink(p, "https://example.com/pay/INV-0042", "example.com/pay")
p = d.add_paragraph(); r = p.add_run("Late payments attract 2% interest per month."); r.italic = True
d.save(f"{OUT}/invoice.docx")

# ---------- proposal.docx ----------
png(f"{OUT}/_pic.png", 40, 24, (60, 140, 90))
d = Document()
d.add_heading("Community Seed Bank Proposal", 0)
d.add_paragraph("Prepared for the Example Foundation. Draft v0.3, October 2026.")
d.add_heading("1. Background", 1)
p = d.add_paragraph("Smallholder farmers in the district lose ")
r = p.add_run("up to 30%"); r.bold = True
p.add_run(" of their seed stock each season. See the ")
add_hyperlink(p, "https://example.com/report", "baseline report")
p.add_run(" and the ")
r = p.add_run("field notes"); r.italic = True
p.add_run(" for detail. The script ")
r = p.add_run("seedctl --report"); r.font.name = "Consolas"
p.add_run(" produces the weekly summary.")
d.add_heading("1.1 Objectives", 2)
for s in ["Reduce seed loss", "Train 40 farmers", "Publish open results"]:
    d.add_paragraph(s, style="List Bullet")
d.add_paragraph("Storage units", style="List Bullet 2")
d.add_paragraph("Hermetic bags", style="List Bullet 2")
d.add_heading("1.2 Steps", 2)
for s in ["Survey sites", "Install storage", "Train farmers"]:
    d.add_paragraph(s, style="List Number")
d.add_heading("2. Budget", 1)
t = d.add_table(rows=3, cols=3); t.style = "Table Grid"
a = t.cell(0, 0).merge(t.cell(0, 2)); a.text = "Budget summary (UGX)"
for i, v in enumerate(["Item", "Quantity", "Cost"]): t.cell(1, i).text = v
for i, v in enumerate(["Hermetic bags", "400", "6,000,000"]): t.cell(2, i).text = v
d.add_heading("3. Risks", 1)
q = d.add_paragraph("Rains may delay site visits by two weeks.", style="Intense Quote")
d.add_heading("4. Appendix", 1)
for line in ["def weekly(loss):", "    return round(loss * 0.3, 2)"]:
    p = d.add_paragraph(); r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9)
d.add_picture(f"{OUT}/_pic.png", width=Cm(3))
d.save(f"{OUT}/proposal.docx")
os.remove(f"{OUT}/_pic.png")

# ---------- tracker.xlsx ----------
wb = Workbook(); ws = wb.active; ws.title = "Q3 Sales"
ws.merge_cells("A1:F1"); ws["A1"] = "Q3 2026 field sales (invented data)"; ws["A1"].font = Font(bold=True, size=14)
hdr = ["Date", "Site", "Farmers", "Revenue (UGX)", "Growth", "Status"]
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=2, column=i, value=h); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEEDD")
sites = ["Mukono", "Wakiso", "Jinja", "Mbale", "Gulu", "Masaka"]
for n in range(12):
    r = 3 + n
    ws.cell(row=r, column=1, value=datetime.date(2026, 7, 1) + datetime.timedelta(days=7 * n)).number_format = "dd/mm/yyyy"
    ws.cell(row=r, column=2, value=sites[n % 6])
    ws.cell(row=r, column=3, value=random.randint(8, 60))
    ws.cell(row=r, column=4, value=random.randint(8, 90) * 100000).number_format = "#,##0"
    ws.cell(row=r, column=5, value=round(random.uniform(-0.1, 0.4), 3)).number_format = "0.0%"
    ws.cell(row=r, column=6, value="done" if n % 3 else "pending")
ws.cell(row=15, column=2, value="Total"); ws.cell(row=15, column=3, value="=SUM(C3:C14)"); ws.cell(row=15, column=4, value="=SUM(D3:D14)").number_format = "#,##0"
ws["B17"] = "Source"; ws["C17"] = "example.com/data"; ws["C17"].hyperlink = "https://example.com/data"
ws2 = wb.create_sheet("Notes"); ws2["A1"] = "Remember"; ws2["A2"] = "Revenue is net of transport."
wb.save(f"{OUT}/tracker.xlsx")

# ---------- training.pptx ----------
png(f"{OUT}/_pic.png", 200, 120, (200, 170, 60))
pr = Presentation()
s = pr.slides.add_slide(pr.slide_layouts[0]); s.shapes.title.text = "Seed Storage Training"; s.placeholders[1].text = "Example Foundation, October 2026"
s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = "Agenda"
tf = s.placeholders[1].text_frame; tf.text = "Why seed is lost"
for txt, lvl in [("Moisture", 1), ("Pests", 1), ("Hermetic storage", 0), ("Costs and returns", 0)]:
    para = tf.add_paragraph(); para.text = txt; para.level = lvl
s.notes_slide.notes_text_frame.text = "Allow 10 minutes. Ask who has lost seed."
s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = "Costs"
tb = s.shapes.add_table(3, 3, Inches(0.5), Inches(1.8), Inches(8), Inches(1.5)).table
for i, v in enumerate(["Item", "Qty", "UGX"]): tb.cell(0, i).text = v
for i, v in enumerate(["Hermetic bag", "400", "6,000,000"]): tb.cell(1, i).text = v
for i, v in enumerate(["Training", "2", "500,000"]): tb.cell(2, i).text = v
s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = "Field photo"
s.shapes.add_picture(f"{OUT}/_pic.png", Inches(1), Inches(2), width=Inches(3))
pr.save(f"{OUT}/training.pptx"); os.remove(f"{OUT}/_pic.png")

# ---------- statement (docx -> pdf) ----------
d = Document()
d.add_heading("Mobile Money Statement", 1)
d.add_paragraph("Account: 0700 000 000    Period: 01 Sep 2026 to 30 Sep 2026    (invented data)")
t = d.add_table(rows=1, cols=6); t.style = "Table Grid"
for i, h in enumerate(["Date", "Transaction ID", "Type", "Details", "Amount (UGX)", "Balance (UGX)"]):
    t.rows[0].cells[i].text = h; t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
bal = 500000
kinds = [("Deposit", 1), ("Payment", -1), ("Transfer", -1), ("Airtime", -1), ("Withdrawal", -1)]
for n in range(ROWS):
    k, sign = kinds[n % 5]; amt = random.randint(5, 150) * 1000; bal += sign * amt
    row = t.add_row().cells
    row[0].text = (datetime.date(2026, 9, 1) + datetime.timedelta(days=n * 30 // ROWS)).strftime("%d/%m/%Y")
    row[1].text = f"MP{26090000 + n * 37}"; row[2].text = k
    row[3].text = random.choice(["Example Shop", "Sample Cooperative", "Test Fuel Station", "Demo Pharmacy"])
    row[4].text = f"{sign * amt:,}"; row[5].text = f"{bal:,}"
for row in t.rows:
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs: r.font.size = Pt(8)
d.save(f"{OUT}/statement.docx")

# ---------- receipt photo ----------
W, H = 560, 760
img = Image.new("RGB", (W, H), (236, 232, 222))
dr = ImageDraw.Draw(img)
try:
    f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", 22)
    fb = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf", 26)
except Exception:
    f = fb = ImageFont.load_default()
lines = [("EXAMPLE AGRO SUPPLIES", fb), ("Plot 12, Example Road", f), ("Tel 0700 000 000", f), ("", f),
         ("RECEIPT No 0731", f), ("Date 06/10/2026", f), ("", f),
         ("Maize seed 25kg x2   170,000", f), ("NPK 50kg x1         160,000", f), ("Transport            40,000", f),
         ("", f), ("TOTAL UGX         370,000", fb), ("Paid by Mobile Money", f), ("Thank you", f)]
y = 40
for text, font in lines:
    dr.text((30, y), text, fill=(30, 30, 30), font=font); y += 40
img = img.rotate(2.2, expand=True, fillcolor=(236, 232, 222)).filter(ImageFilter.GaussianBlur(0.6))
px = img.load()
for _ in range(9000):
    x, yy = random.randrange(img.width), random.randrange(img.height); v = random.randint(-25, 25)
    r, g, b = px[x, yy]; px[x, yy] = (max(0, min(255, r + v)), max(0, min(255, g + v)), max(0, min(255, b + v)))
img.save(f"{OUT}/receipt.png"); img.convert("RGB").save(f"{OUT}/receipt.jpg", quality=70)

# ---------- html / xml / text ----------
html = """<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>Seed bank update</title></head><body>
<h1>Seed bank update</h1>
<p>We reached <strong>three</strong> sites and <em>two</em> more are planned. Read the <a href="https://example.com/report">full report</a>.</p>
<h2>Progress</h2>
<table><thead><tr><th>Site</th><th>Farmers</th><th>Status</th></tr></thead>
<tbody><tr><td>Mukono</td><td>42</td><td>done</td></tr><tr><td>Wakiso</td><td>17</td><td>pending</td></tr></tbody></table>
<ul><li>Survey</li><li>Install<ul><li>Bags</li><li>Racks</li></ul></li></ul>
<ol><li>First</li><li>Second</li></ol>
<blockquote>Seed saved is food saved.</blockquote>
<pre><code class="language-python">print("hello")
</code></pre>
<p>Use <code>seedctl</code> daily. Fish &amp; chips &lt;ok&gt; caf&eacute;.</p>
<script>alert(1)</script><style>p{color:red}</style>
</body></html>"""
open(f"{OUT}/page.html", "w", encoding="utf-8").write(html)
open(f"{OUT}/page-latin1.html", "w", encoding="cp1252").write(html.replace('<meta charset="utf-8">', '<meta http-equiv="Content-Type" content="text/html; charset=windows-1252">').replace("Seed bank update</h1>", "Seed bank update – café</h1>"))
open(f"{OUT}/page-utf16.html", "wb").write(b"\xff\xfe" + html.encode("utf-16-le"))
open(f"{OUT}/records.xml", "w").write('<?xml version="1.0" encoding="UTF-8"?>\n<visits>\n  <visit site="Mukono" farmers="42"><note>Done on time</note></visit>\n  <visit site="Wakiso" farmers="17"><note>Rain delay</note></visit>\n  <visit site="Jinja" farmers="23"><note>Pending</note></visit>\n</visits>\n')
open(f"{OUT}/notes.xml", "w").write('<?xml version="1.0"?>\n<report><title>Field report</title><section><heading>Findings</heading><para>Seed loss fell to <b>12%</b>.</para><para>Training reached 40 farmers.</para></section></report>\n')
open(f"{OUT}/data.csv", "w").write("Site,Farmers,Status\nMukono,42,done\nWakiso,17,pending\n")
open(f"{OUT}/data.tsv", "w").write("Site\tFarmers\tStatus\nMukono\t42\tdone\nWakiso\t17\tpending\n")
open(f"{OUT}/plain.txt", "w").write("Field note\n\nMukono: 42 farmers trained.\nWakiso: 17 pending.\n")
print("generated", sorted(os.listdir(OUT)))
